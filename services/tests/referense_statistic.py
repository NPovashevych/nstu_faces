import logging
from datetime import datetime
from itertools import combinations
from pathlib import Path

import cv2
import numpy as np
from insightface.app import FaceAnalysis
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services.config import PERSONS_FOLDER, REFERENSE_STATISTIC


MODEL_NAME = "buffalo_l"
DET_SIZE = (640, 640)
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}

BLUR_THRESHOLDS = [5, 10, 15, 20, 30, 40, 50]
BBOX_THRESHOLDS = [30, 40, 50, 60, 80, 100, 120]
DISTANCE_THRESHOLDS = [0.25, 0.35, 0.42, 0.50, 0.60, 0.65, 0.72, 0.80]
DET_SCORE_THRESHOLDS = [0.50, 0.60, 0.70, 0.80, 0.90]

LOG_FILE = Path("../logs/reference_statistic.log")


def setup_logger():
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("reference_statistic")
    logger.setLevel(logging.INFO)
    logger.propagate = False

    if logger.handlers:
        logger.handlers.clear()

    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s", datefmt="%Y-%m-%d %H:%M:%S")

    file_handler = logging.FileHandler(LOG_FILE, mode="w", encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)

    console_handler = logging.StreamHandler()
    console_handler.setFormatter(formatter)
    logger.addHandler(console_handler)

    return logger


logger = setup_logger()


def read_image(file_path):
    image_data = np.fromfile(file_path, dtype=np.uint8)

    if image_data.size == 0:
        return None

    return cv2.imdecode(image_data, cv2.IMREAD_COLOR)


def normalize_embedding(embedding):
    embedding = np.asarray(embedding, dtype=np.float32)
    norm = np.linalg.norm(embedding)
    if norm == 0:
        return None
    return embedding / norm


def cosine_distance(embedding_1, embedding_2):
    similarity = float(np.dot(embedding_1, embedding_2))
    similarity = float(np.clip(similarity, -1.0, 1.0))
    return 1.0 - similarity


def calculate_blur(image):
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    return float(cv2.Laplacian(gray, cv2.CV_64F).var())


def safe_percentile(values, percentile):
    return None if not values else float(np.percentile(values, percentile))


def safe_mean(values):
    return None if not values else float(np.mean(values))


def safe_median(values):
    return None if not values else float(np.median(values))


def safe_min(values):
    return None if not values else float(min(values))


def safe_max(values):
    return None if not values else float(max(values))


def percent(count, total):
    return 0.0 if not total else count / total * 100.0


def get_largest_face(faces):
    return max(faces, key=lambda face: float(face.bbox[2] - face.bbox[0]) * float(face.bbox[3] - face.bbox[1]))


def clean_excel_value(value):
    if value is None:
        return None
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    return value


def style_header(sheet, fill, border, row_number=1):
    for cell in sheet[row_number]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def auto_width(sheet, max_width=45):
    widths = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))

    for column_index, width in widths.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = min(width + 2, max_width)


def write_dict_sheet(sheet, rows, header_fill, thin_border):
    if not rows:
        return

    headers = list(rows[0].keys())
    sheet.append(headers)

    for row in rows:
        sheet.append([clean_excel_value(row.get(header)) for header in headers])

    style_header(sheet, header_fill, thin_border)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    auto_width(sheet)


def add_distribution(sheet, start_row, title, values, section_fill, header_fill):
    sheet.cell(row=start_row, column=1, value=title)
    sheet.cell(row=start_row, column=1).font = Font(bold=True)
    sheet.cell(row=start_row, column=1).fill = section_fill

    headers = ["N", "MIN", "P1", "P5", "P10", "P25", "MEDIAN", "P75", "P90", "P95", "P99", "MAX", "MEAN", "STD"]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row + 1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    if values:
        stats = [len(values), safe_min(values), safe_percentile(values, 1), safe_percentile(values, 5), safe_percentile(values, 10), safe_percentile(values, 25), safe_percentile(values, 50), safe_percentile(values, 75), safe_percentile(values, 90), safe_percentile(values, 95), safe_percentile(values, 99), safe_max(values), safe_mean(values), float(np.std(values))]
    else:
        stats = [0] + [None] * (len(headers) - 1)

    for column, value in enumerate(stats, start=1):
        sheet.cell(row=start_row + 2, column=column, value=clean_excel_value(value))

    return start_row + 4


def add_threshold_section(sheet, start_row, title, values, thresholds, mode, warning_fill, header_fill):
    sheet.cell(row=start_row, column=1, value=title)
    sheet.cell(row=start_row, column=1).font = Font(bold=True)
    sheet.cell(row=start_row, column=1).fill = warning_fill

    for column, header in enumerate(["Поріг", "Кількість", "%"], start=1):
        cell = sheet.cell(row=start_row + 1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    current_row = start_row + 2
    for threshold in thresholds:
        if mode == "below":
            count = sum(value < threshold for value in values)
            label = f"< {threshold}"
        else:
            count = sum(value > threshold for value in values)
            label = f"> {threshold}"

        sheet.cell(row=current_row, column=1, value=label)
        sheet.cell(row=current_row, column=2, value=count)
        sheet.cell(row=current_row, column=3, value=percent(count, len(values)))
        current_row += 1

    return current_row + 2


def main():
    logger.info("=" * 80)
    logger.info("ЗАВАНТАЖЕННЯ INSIGHTFACE")
    logger.info("=" * 80)

    model = FaceAnalysis(name=MODEL_NAME, providers=["CUDAExecutionProvider", "CPUExecutionProvider"])
    model.prepare(ctx_id=0, det_size=DET_SIZE)
    logger.info("InsightFace готовий")

    if not PERSONS_FOLDER.exists():
        raise FileNotFoundError(f"Папку еталонів не знайдено: {PERSONS_FOLDER}")

    person_folders = sorted([folder for folder in PERSONS_FOLDER.iterdir() if folder.is_dir()], key=lambda path: path.name.lower())
    total_persons = len(person_folders)
    total_images = sum(1 for person_folder in person_folders for file_path in person_folder.iterdir() if file_path.is_file() and file_path.suffix.lower() in IMAGE_EXTENSIONS)

    logger.info("=" * 80)
    logger.info("ЕТАЛОННА КОЛЕКЦІЯ")
    logger.info("=" * 80)
    logger.info("Папка: %s", PERSONS_FOLDER)
    logger.info("Персон: %s", total_persons)
    logger.info("Фото дозволених форматів: %s", total_images)
    logger.info("Дозволені формати: %s", ", ".join(sorted(IMAGE_EXTENSIONS)))

    photo_rows = []
    pair_rows = []
    person_rows = []
    processed_images = 0

    for person_index, person_folder in enumerate(person_folders, start=1):
        image_files = sorted([file_path for file_path in person_folder.iterdir() if file_path.is_file() and file_path.suffix.lower() in IMAGE_EXTENSIONS], key=lambda path: path.name.lower())
        logger.info("PERSON %s/%s | %s | фото: %s", person_index, total_persons, person_folder.name, len(image_files))

        person_photo_rows = []
        person_embeddings = []
        unreadable_count = 0
        no_face_count = 0
        one_face_count = 0
        multiple_faces_count = 0
        insightface_error_count = 0

        for file_path in image_files:
            processed_images += 1
            logger.info("PHOTO %s/%s | %s | %s", processed_images, total_images, person_folder.name, file_path.name)

            row = {
                "person_folder": person_folder.name,
                "file_name": file_path.name,
                "file_path": str(file_path),
                "image_width": None,
                "image_height": None,
                "image_area": None,
                "image_min_side": None,
                "blur_value": None,
                "faces_count": None,
                "bbox_x1": None,
                "bbox_y1": None,
                "bbox_x2": None,
                "bbox_y2": None,
                "bbox_width": None,
                "bbox_height": None,
                "bbox_area": None,
                "bbox_min_side": None,
                "bbox_area_ratio": None,
                "bbox_width_ratio": None,
                "bbox_height_ratio": None,
                "det_score": None,
                "distance_to_person_mean": None,
                "mean_distance_to_other_photos": None,
                "max_distance_to_other_photo": None,
                "status": None,
            }

            image = read_image(file_path)
            if image is None:
                unreadable_count += 1
                row["status"] = "unreadable"
                photo_rows.append(row)
                person_photo_rows.append(row)
                logger.warning("Файл не читається: %s", file_path)
                continue

            image_height, image_width = image.shape[:2]
            image_area = image_width * image_height
            image_min_side = min(image_width, image_height)

            row["image_width"] = image_width
            row["image_height"] = image_height
            row["image_area"] = image_area
            row["image_min_side"] = image_min_side

            blur_value = calculate_blur(image)
            row["blur_value"] = blur_value

            try:
                faces = model.get(image)
            except Exception:
                insightface_error_count += 1
                row["status"] = "insightface_error"
                photo_rows.append(row)
                person_photo_rows.append(row)
                logger.exception("InsightFace error: %s", file_path)
                continue

            faces_count = len(faces)
            row["faces_count"] = faces_count

            if faces_count == 0:
                no_face_count += 1
                row["status"] = "no_face"
                photo_rows.append(row)
                person_photo_rows.append(row)
                logger.warning("Обличчя не знайдено | blur=%.2f | %s", blur_value, file_path)
                continue

            if faces_count == 1:
                one_face_count += 1
                row["status"] = "ok"
            else:
                multiple_faces_count += 1
                row["status"] = "multiple_faces"
                logger.warning("Знайдено %s облич | %s", faces_count, file_path)

            face = get_largest_face(faces)
            x1, y1, x2, y2 = [float(value) for value in face.bbox]
            bbox_width = x2 - x1
            bbox_height = y2 - y1
            bbox_area = bbox_width * bbox_height
            bbox_min_side = min(bbox_width, bbox_height)

            row["bbox_x1"] = x1
            row["bbox_y1"] = y1
            row["bbox_x2"] = x2
            row["bbox_y2"] = y2
            row["bbox_width"] = bbox_width
            row["bbox_height"] = bbox_height
            row["bbox_area"] = bbox_area
            row["bbox_min_side"] = bbox_min_side
            row["bbox_area_ratio"] = bbox_area / image_area if image_area > 0 else None
            row["bbox_width_ratio"] = bbox_width / image_width if image_width > 0 else None
            row["bbox_height_ratio"] = bbox_height / image_height if image_height > 0 else None

            det_score = getattr(face, "det_score", None)
            row["det_score"] = float(det_score) if det_score is not None else None

            embedding = getattr(face, "embedding", None)
            if embedding is not None:
                normalized_embedding = normalize_embedding(embedding)
                if normalized_embedding is not None:
                    person_embeddings.append({"row": row, "file_name": file_path.name, "embedding": normalized_embedding})

            photo_rows.append(row)
            person_photo_rows.append(row)

            if row["det_score"] is not None:
                logger.info("RESULT | %sx%s | blur=%.2f | faces=%s | bbox=%.1fx%.1f | det=%.4f", image_width, image_height, blur_value, faces_count, bbox_width, bbox_height, row["det_score"])
            else:
                logger.info("RESULT | %sx%s | blur=%.2f | faces=%s | bbox=%.1fx%.1f | det=None", image_width, image_height, blur_value, faces_count, bbox_width, bbox_height)

        person_pair_distances = []

        if person_embeddings:
            embedding_matrix = np.vstack([item["embedding"] for item in person_embeddings])
            mean_embedding = normalize_embedding(np.mean(embedding_matrix, axis=0))

            if mean_embedding is not None:
                for item in person_embeddings:
                    item["row"]["distance_to_person_mean"] = cosine_distance(item["embedding"], mean_embedding)

        distances_by_file = {item["file_name"]: [] for item in person_embeddings}

        for item_1, item_2 in combinations(person_embeddings, 2):
            similarity = float(np.dot(item_1["embedding"], item_2["embedding"]))
            similarity = float(np.clip(similarity, -1.0, 1.0))
            distance = 1.0 - similarity

            person_pair_distances.append(distance)
            distances_by_file[item_1["file_name"]].append(distance)
            distances_by_file[item_2["file_name"]].append(distance)

            pair_rows.append({
                "person_folder": person_folder.name,
                "file_1": item_1["file_name"],
                "file_2": item_2["file_name"],
                "cosine_similarity": similarity,
                "cosine_distance": distance,
                "over_0_25": distance > 0.25,
                "over_0_35": distance > 0.35,
                "over_0_42": distance > 0.42,
                "over_0_50": distance > 0.50,
                "over_0_60": distance > 0.60,
                "over_0_65": distance > 0.65,
                "over_0_72": distance > 0.72,
                "over_0_80": distance > 0.80,
            })

        for item in person_embeddings:
            file_distances = distances_by_file[item["file_name"]]
            if file_distances:
                item["row"]["mean_distance_to_other_photos"] = float(np.mean(file_distances))
                item["row"]["max_distance_to_other_photo"] = float(max(file_distances))

        blur_values = [row["blur_value"] for row in person_photo_rows if row["blur_value"] is not None]
        bbox_values = [row["bbox_min_side"] for row in person_photo_rows if row["bbox_min_side"] is not None]
        det_score_values = [row["det_score"] for row in person_photo_rows if row["det_score"] is not None]
        distance_to_mean_values = [row["distance_to_person_mean"] for row in person_photo_rows if row["distance_to_person_mean"] is not None]
        pairs_count = len(person_pair_distances)

        person_summary = {
            "person_folder": person_folder.name,
            "photos_count": len(image_files),
            "readable_count": len(image_files) - unreadable_count,
            "unreadable_count": unreadable_count,
            "no_face_count": no_face_count,
            "one_face_count": one_face_count,
            "multiple_faces_count": multiple_faces_count,
            "insightface_error_count": insightface_error_count,
            "embeddings_count": len(person_embeddings),
            "min_blur": safe_min(blur_values),
            "median_blur": safe_median(blur_values),
            "mean_blur": safe_mean(blur_values),
            "min_bbox_side": safe_min(bbox_values),
            "median_bbox_side": safe_median(bbox_values),
            "mean_bbox_side": safe_mean(bbox_values),
            "min_det_score": safe_min(det_score_values),
            "median_det_score": safe_median(det_score_values),
            "mean_det_score": safe_mean(det_score_values),
            "max_distance_to_mean": safe_max(distance_to_mean_values),
            "median_distance_to_mean": safe_median(distance_to_mean_values),
            "mean_distance_to_mean": safe_mean(distance_to_mean_values),
            "pairs_count": pairs_count,
            "min_pairwise_distance": safe_min(person_pair_distances),
            "median_pairwise_distance": safe_median(person_pair_distances),
            "mean_pairwise_distance": safe_mean(person_pair_distances),
            "max_pairwise_distance": safe_max(person_pair_distances),
        }

        for threshold in BLUR_THRESHOLDS:
            count = sum(value < threshold for value in blur_values)
            person_summary[f"blur_below_{threshold}_count"] = count
            person_summary[f"blur_below_{threshold}_percent"] = percent(count, len(blur_values))

        for threshold in BBOX_THRESHOLDS:
            count = sum(value < threshold for value in bbox_values)
            person_summary[f"bbox_below_{threshold}_count"] = count
            person_summary[f"bbox_below_{threshold}_percent"] = percent(count, len(bbox_values))

        for threshold in DISTANCE_THRESHOLDS:
            count = sum(value > threshold for value in person_pair_distances)
            threshold_name = str(threshold).replace(".", "_")
            person_summary[f"pairs_over_{threshold_name}_count"] = count
            person_summary[f"pairs_over_{threshold_name}_percent"] = percent(count, pairs_count)

        person_rows.append(person_summary)

    all_blur = [row["blur_value"] for row in photo_rows if row["blur_value"] is not None]
    all_image_width = [row["image_width"] for row in photo_rows if row["image_width"] is not None]
    all_image_height = [row["image_height"] for row in photo_rows if row["image_height"] is not None]
    all_image_area = [row["image_area"] for row in photo_rows if row["image_area"] is not None]
    all_image_min_side = [row["image_min_side"] for row in photo_rows if row["image_min_side"] is not None]
    all_bbox_width = [row["bbox_width"] for row in photo_rows if row["bbox_width"] is not None]
    all_bbox_height = [row["bbox_height"] for row in photo_rows if row["bbox_height"] is not None]
    all_bbox_area = [row["bbox_area"] for row in photo_rows if row["bbox_area"] is not None]
    all_bbox_min_side = [row["bbox_min_side"] for row in photo_rows if row["bbox_min_side"] is not None]
    all_bbox_area_ratio = [row["bbox_area_ratio"] for row in photo_rows if row["bbox_area_ratio"] is not None]
    all_det_score = [row["det_score"] for row in photo_rows if row["det_score"] is not None]
    all_distance_to_mean = [row["distance_to_person_mean"] for row in photo_rows if row["distance_to_person_mean"] is not None]
    all_pairwise_distance = [row["cosine_distance"] for row in pair_rows]

    logger.info("=" * 80)
    logger.info("ФОРМУВАННЯ EXCEL")
    logger.info("=" * 80)

    workbook = Workbook()
    title_sheet = workbook.active
    title_sheet.title = "Титул"
    photo_sheet = workbook.create_sheet("Фото")
    pair_sheet = workbook.create_sheet("Пари")
    person_sheet = workbook.create_sheet("Персони")

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="E2F0D9")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    title_fill = PatternFill("solid", fgColor="BDD7EE")
    thin_border = Border(bottom=Side(style="thin", color="D9D9D9"))

    write_dict_sheet(photo_sheet, photo_rows, header_fill, thin_border)
    write_dict_sheet(pair_sheet, pair_rows, header_fill, thin_border)
    write_dict_sheet(person_sheet, person_rows, header_fill, thin_border)

    title_sheet["A1"] = "REFERENCE FACE QUALITY BASELINE"
    title_sheet["A1"].font = Font(bold=True, size=16)
    title_sheet["A1"].fill = title_fill
    title_sheet["A3"] = "Дата аналізу"
    title_sheet["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    title_sheet["A4"] = "Каталог"
    title_sheet["B4"] = str(PERSONS_FOLDER)
    title_sheet["A5"] = "InsightFace model"
    title_sheet["B5"] = MODEL_NAME
    title_sheet["A6"] = "det_size"
    title_sheet["B6"] = f"{DET_SIZE[0]}x{DET_SIZE[1]}"
    title_sheet["A7"] = "Дозволені формати"
    title_sheet["B7"] = ", ".join(sorted(IMAGE_EXTENSIONS))
    title_sheet["A8"] = "Blur metric"
    title_sheet["B8"] = "Variance of Laplacian, grayscale"
    title_sheet["A9"] = "Cosine distance"
    title_sheet["B9"] = "1 - cosine similarity для L2-normalized embeddings"

    row_number = 11
    title_sheet.cell(row=row_number, column=1, value="ЗАГАЛЬНА СТАТИСТИКА")
    title_sheet.cell(row=row_number, column=1).font = Font(bold=True)
    title_sheet.cell(row=row_number, column=1).fill = section_fill

    summary_values = [
        ("Папок / персон", len(person_rows)),
        ("Фото", len(photo_rows)),
        ("Unreadable", sum(row["status"] == "unreadable" for row in photo_rows)),
        ("InsightFace errors", sum(row["status"] == "insightface_error" for row in photo_rows)),
        ("0 faces", sum(row["faces_count"] == 0 for row in photo_rows)),
        ("1 face", sum(row["faces_count"] == 1 for row in photo_rows)),
        ("2+ faces", sum(row["faces_count"] is not None and row["faces_count"] > 1 for row in photo_rows)),
        ("Фото з embedding", len(all_distance_to_mean)),
        ("Pairwise comparisons", len(pair_rows)),
    ]

    row_number += 1
    for name, value in summary_values:
        title_sheet.cell(row=row_number, column=1, value=name)
        title_sheet.cell(row=row_number, column=2, value=value)
        row_number += 1

    row_number += 2
    distributions = [
        ("РОЗПОДІЛ: ширина фото", all_image_width),
        ("РОЗПОДІЛ: висота фото", all_image_height),
        ("РОЗПОДІЛ: площа фото", all_image_area),
        ("РОЗПОДІЛ: min side фото", all_image_min_side),
        ("РОЗПОДІЛ: blur", all_blur),
        ("РОЗПОДІЛ: bbox width", all_bbox_width),
        ("РОЗПОДІЛ: bbox height", all_bbox_height),
        ("РОЗПОДІЛ: bbox area", all_bbox_area),
        ("РОЗПОДІЛ: bbox min side", all_bbox_min_side),
        ("РОЗПОДІЛ: bbox / image area", all_bbox_area_ratio),
        ("РОЗПОДІЛ: det_score", all_det_score),
        ("РОЗПОДІЛ: distance to person mean", all_distance_to_mean),
        ("РОЗПОДІЛ: pairwise cosine distance", all_pairwise_distance),
    ]

    for title, values in distributions:
        row_number = add_distribution(title_sheet, row_number, title, values, section_fill, header_fill)

    row_number = add_threshold_section(title_sheet, row_number, "BLUR — частка нижче порогу", all_blur, BLUR_THRESHOLDS, "below", warning_fill, header_fill)
    row_number = add_threshold_section(title_sheet, row_number, "BBOX MIN SIDE — частка нижче порогу", all_bbox_min_side, BBOX_THRESHOLDS, "below", warning_fill, header_fill)
    row_number = add_threshold_section(title_sheet, row_number, "PAIRWISE DISTANCE — частка вище порогу", all_pairwise_distance, DISTANCE_THRESHOLDS, "above", warning_fill, header_fill)
    row_number = add_threshold_section(title_sheet, row_number, "DET SCORE — частка нижче порогу", all_det_score, DET_SCORE_THRESHOLDS, "below", warning_fill, header_fill)

    title_sheet.cell(row=row_number, column=1, value="РОЗПОДІЛ КІЛЬКОСТІ ОБЛИЧ")
    title_sheet.cell(row=row_number, column=1).font = Font(bold=True)
    title_sheet.cell(row=row_number, column=1).fill = section_fill
    row_number += 1

    faces_categories = [
        ("0 faces", sum(row["faces_count"] == 0 for row in photo_rows)),
        ("1 face", sum(row["faces_count"] == 1 for row in photo_rows)),
        ("2 faces", sum(row["faces_count"] == 2 for row in photo_rows)),
        ("3 faces", sum(row["faces_count"] == 3 for row in photo_rows)),
        ("4+ faces", sum(row["faces_count"] is not None and row["faces_count"] >= 4 for row in photo_rows)),
    ]

    valid_face_count_rows = sum(row["faces_count"] is not None for row in photo_rows)
    for label, count in faces_categories:
        title_sheet.cell(row=row_number, column=1, value=label)
        title_sheet.cell(row=row_number, column=2, value=count)
        title_sheet.cell(row=row_number, column=3, value=percent(count, valid_face_count_rows))
        row_number += 1

    title_sheet.column_dimensions["A"].width = 38
    for column in range(2, 15):
        title_sheet.column_dimensions[get_column_letter(column)].width = 14
    title_sheet.freeze_panes = "A3"

    for row in title_sheet.iter_rows():
        for cell in row:
            if cell.column == 3 and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"

    for sheet in [photo_sheet, pair_sheet, person_sheet]:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"

    REFERENSE_STATISTIC.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(REFERENSE_STATISTIC)

    logger.info("=" * 80)
    logger.info("ГОТОВО")
    logger.info("=" * 80)
    logger.info("Звіт: %s", REFERENSE_STATISTIC)
    logger.info("Лог: %s", LOG_FILE.resolve())
    logger.info("Персон: %s", len(person_rows))
    logger.info("Фото: %s", len(photo_rows))
    logger.info("Пар фото: %s", len(pair_rows))
    logger.info("Листи: Титул | Фото | Пари | Персони")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        logger.exception("Критична помилка під час формування статистики")
        raise
